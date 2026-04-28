from __future__ import annotations

import base64
import io
import unittest
import zipfile
from xml.etree import ElementTree

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from checker.validator import Validator


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgG"
    "M9t6RpwAAAABJRU5ErkJggg=="
)


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _set_docx_page_count(content: bytes, pages: int) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content), "r") as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/app.xml":
                root = ElementTree.fromstring(content)
                for element in root.iter():
                    if _local_name(element.tag) == "Pages":
                        element.text = str(pages)
                        break
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, content)
    return output.getvalue()


def _configure_valid_section(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)


def _add_body_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return paragraph


def _docx_bytes(doc: Document, *, pages: int) -> io.BytesIO:
    output = io.BytesIO()
    doc.save(output)
    output = io.BytesIO(_set_docx_page_count(output.getvalue(), pages))
    output.seek(0)
    return output


def _valid_docx_bytes() -> io.BytesIO:
    doc = Document()
    _configure_valid_section(doc.sections[0])
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    _add_body_paragraph(doc, "Текст статьи " * 620)
    return _docx_bytes(doc, pages=4)


def _validate_global(content: io.BytesIO):
    validator = Validator(content)
    validator.validate_global_requirements()
    return validator.errors, validator.errors_eng


def _add_superscript_run(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.superscript = True
    run.bold = True


def _add_article_paragraph(
        doc: Document,
        text: str,
        *,
        bold: bool = False,
        italic: bool = False,
        alignment=None,
        first_line_indent_cm: float | None = None,
):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    if alignment is not None:
        paragraph.alignment = alignment
    if first_line_indent_cm is not None:
        paragraph.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return paragraph


def _add_labeled_article_paragraph(doc: Document, label: str, body: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    label_run = paragraph.add_run(label)
    label_run.font.name = "Times New Roman"
    label_run.font.size = Pt(12)
    label_run.bold = True
    body_run = paragraph.add_run(" " + body)
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(12)
    return paragraph


def _add_small_picture(doc: Document) -> None:
    doc.add_picture(io.BytesIO(PNG_1X1), width=Cm(1))


def _add_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    for index, cell in enumerate(table.rows[0].cells, 1):
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(str(index))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def _add_table_caption_and_title(doc: Document, number: int, title: str = "Заголовок таблицы") -> None:
    _add_article_paragraph(doc, f"Таблица {number}", italic=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_article_paragraph(doc, title, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)


def _add_author_paragraph(doc: Document, name: str, marker: str, email: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    name_run = paragraph.add_run(name)
    name_run.font.name = "Times New Roman"
    name_run.font.size = Pt(12)
    name_run.bold = True
    _add_superscript_run(paragraph, marker)
    email_run = paragraph.add_run(f" {email}")
    email_run.font.name = "Times New Roman"
    email_run.font.size = Pt(12)
    return paragraph


def _add_reference_paragraph(doc: Document, number: int, text: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Cm(1)
    paragraph.paragraph_format.first_line_indent = Cm(-1)

    prefix_run = paragraph.add_run(f"[{number}] ")
    prefix_run.font.name = "Times New Roman"
    prefix_run.font.size = Pt(12)
    body_run = paragraph.add_run("\t" + (text or f"Source title {number}. Moscow, Publisher, 2024, 10 p."))
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(12)
    return paragraph


def _long_text(seed: str, min_length: int = 700) -> str:
    parts = []
    while len(" ".join(parts)) < min_length:
        parts.append(seed)
    return " ".join(parts)


def _valid_article_document() -> Document:
    doc = Document()
    _configure_valid_section(doc.sections[0])
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    _add_article_paragraph(doc, "УДК 520.607")
    _add_article_paragraph(doc, "Экономика многоразовости будущей космонавтики и вопросы оперативности", bold=True)
    _add_author_paragraph(doc, "Иванов Иван Иванович", "1", "ivanov@example.com")
    _add_article_paragraph(doc, "SPIN-код: 1234-5678")
    _add_author_paragraph(doc, "Петров Петр Петрович", "2(*)", "petrov@example.com")
    _add_article_paragraph(doc, "1 ФГБУ НПО Тайфун, Москва, Россия", italic=True)
    _add_article_paragraph(doc, "2 МГТУ им. Н.Э. Баумана, Москва, Россия", italic=True)
    _add_labeled_article_paragraph(
        doc,
        "Аннотация.",
        _long_text("Рассмотрены перспективы развития технологий и экономики многоразовых космических систем."),
    )
    _add_labeled_article_paragraph(
        doc,
        "Ключевые слова:",
        "энергообеспечение, загрязнение окружающей среды, нанотехнологии, экологическое равновесие, свойства наноматериалов",
    )
    _add_article_paragraph(doc, "Economy of Reusability of Future Cosmonautics and Issues of Efficiency", bold=True)
    _add_author_paragraph(doc, "Ivanov Ivan Ivanovich", "1", "ivanov@example.com")
    _add_article_paragraph(doc, "SPIN-code: 1234-5678")
    _add_author_paragraph(doc, "Petrov Petr Petrovich", "2(*)", "petrov@example.com")
    _add_article_paragraph(doc, "1 FSBI NPO Typhoon, Moscow, Russia", italic=True)
    _add_article_paragraph(doc, "2 BMSTU, Moscow, Russia", italic=True)
    _add_labeled_article_paragraph(
        doc,
        "Abstract.",
        _long_text("The article considers technological and economic aspects of reusable space systems."),
    )
    _add_labeled_article_paragraph(
        doc,
        "Keywords:",
        "energy supply, environmental pollution, nanotechnology, ecological balance, material properties",
    )
    _add_article_paragraph(
        doc,
        "Introduction. "
        + ("This main article text is intentionally written in English. " * 80)
        + "The first source is cited with a page number [1, p. 17]. "
        + "The remaining sources are cited as a range [2-5].",
    )
    _add_article_paragraph(doc, "References", bold=True)
    for number in range(1, 6):
        _add_reference_paragraph(doc, number)
    return doc


def _valid_article_docx_bytes() -> io.BytesIO:
    return _docx_bytes(_valid_article_document(), pages=4)


def _main_text_paragraph(doc: Document):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Introduction."):
            return paragraph
    raise AssertionError("Main text paragraph was not found")


def _reference_paragraphs(doc: Document):
    return [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("[")]


def _insert_main_text_paragraph_before_references(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text == "References":
            inserted = paragraph.insert_paragraph_before()
            inserted.paragraph_format.line_spacing = 1.5
            run = inserted.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            return inserted
    raise AssertionError("References heading was not found")


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class ValidatorGlobalRequirementsTest(unittest.TestCase):
    def test_global_requirements_pass_for_valid_docx(self):
        errors_ru, errors_en = _validate_global(_valid_docx_bytes())

        self.assertEqual([], errors_ru)
        self.assertEqual([], errors_en)

    def test_global_requirements_report_common_violations(self):
        doc = Document()
        section = doc.sections[0]
        section.page_width = Cm(20)
        section.page_height = Cm(28)
        section.top_margin = Cm(1)
        section.bottom_margin = Cm(1)
        section.left_margin = Cm(1)
        section.right_margin = Cm(1)

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run("Short text")
        run.font.name = "Arial"
        run.font.size = Pt(14)
        _add_hyperlink(paragraph, "link", "https://example.com")
        doc.add_section()

        errors_ru, _ = _validate_global(_docx_bytes(doc, pages=2))

        self.assertIn("Материалы должны быть представлены на формате А4", errors_ru)
        self.assertIn("Отступ сверху должен равняться 2 см", errors_ru)
        self.assertIn("Общий объем статьи должен составлять 4-6 страниц", errors_ru)
        self.assertIn("Минимальный объем статьи должен составлять 6000 знаков с пробелами", errors_ru)
        self.assertIn(
            "В тексте статьи необходимо использовать шрифт Times New Roman, за исключением математических формул",
            errors_ru,
        )
        self.assertIn("Межстрочный интервал должен равняться 1.5", errors_ru)
        self.assertIn("Применять гиперссылки в тексте не допускается", errors_ru)
        self.assertIn("Разрывы разделов внутри текста не допускаются", errors_ru)

    def test_font_check_ignores_empty_runs_and_formula_runs(self):
        doc = Document()
        _configure_valid_section(doc.sections[0])
        _add_body_paragraph(doc, "Текст статьи " * 620)
        paragraph = _add_body_paragraph(doc, "Формула: ")

        empty_run = paragraph.add_run("")
        empty_run.font.name = "Arial"
        empty_run.font.size = Pt(18)

        formula_run = paragraph.add_run("x")
        formula_run.font.name = "Cambria Math"
        formula_run.font.size = Pt(10)
        formula_run._element.append(
            parse_xml(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                "<m:r><m:t>x</m:t></m:r>"
                "</m:oMath>"
            )
        )

        errors_ru, _ = _validate_global(_docx_bytes(doc, pages=4))

        self.assertNotIn(
            "В тексте статьи необходимо использовать шрифт Times New Roman, за исключением математических формул",
            errors_ru,
        )
        self.assertNotIn("Размер шрифта должен быть 12", errors_ru)

    def test_text_format_checks_ignore_object_only_and_trailing_empty_paragraphs(self):
        doc = Document()
        _configure_valid_section(doc.sections[0])
        _add_body_paragraph(doc, "Текст статьи " * 620)

        _add_small_picture(doc)
        picture_paragraph = doc.paragraphs[-1]
        picture_paragraph.paragraph_format.line_spacing = 1

        formula_paragraph = doc.add_paragraph()
        formula_paragraph.paragraph_format.line_spacing = 1
        formula_run = formula_paragraph.add_run("x")
        formula_run.font.name = "Cambria Math"
        formula_run.font.size = Pt(10)
        formula_run._element.append(
            parse_xml(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                "<m:r><m:t>x</m:t></m:r>"
                "</m:oMath>"
            )
        )

        table = doc.add_table(rows=1, cols=1)
        table_cell_paragraph = table.rows[0].cells[0].paragraphs[0]
        table_cell_paragraph.paragraph_format.line_spacing = 1

        trailing_empty = doc.add_paragraph()
        trailing_empty.paragraph_format.line_spacing = 1
        empty_run = trailing_empty.add_run("")
        empty_run.font.name = "Arial"
        empty_run.font.size = Pt(18)

        errors_ru, _ = _validate_global(_docx_bytes(doc, pages=4))

        self.assertNotIn(
            "В тексте статьи необходимо использовать шрифт Times New Roman, за исключением математических формул",
            errors_ru,
        )
        self.assertNotIn("Размер шрифта должен быть 12", errors_ru)
        self.assertNotIn("Межстрочный интервал должен равняться 1.5", errors_ru)

    def test_hyperlink_check_allows_email_links(self):
        doc = Document()
        _configure_valid_section(doc.sections[0])
        _add_body_paragraph(doc, "Текст статьи " * 620)
        paragraph = _add_body_paragraph(doc, "Контакт: ")
        _add_hyperlink(paragraph, "e-mail", "mailto:author@example.com")

        errors_ru, _ = _validate_global(_docx_bytes(doc, pages=4))

        self.assertNotIn("Применять гиперссылки в тексте не допускается", errors_ru)

    def test_valid_article_structure_metadata_and_annotations_pass(self):
        errors_ru, errors_en = Validator(_valid_article_docx_bytes()).validate()

        self.assertEqual([], errors_ru)
        self.assertEqual([], errors_en)

    def test_article_structure_requires_main_text(self):
        doc = _valid_article_document()
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith("Introduction."):
                paragraph.clear()
                break

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Основной текст статьи не найден", errors_ru)

    def test_article_structure_requires_sources_heading(self):
        doc = _valid_article_document()
        for paragraph in doc.paragraphs:
            if paragraph.text == "References":
                paragraph.clear()
                break

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Заголовок списка источников должен быть References", errors_ru)

    def test_metadata_and_keyword_violations_are_reported(self):
        doc = _valid_article_document()
        doc.paragraphs[0].text = "УДК wrong"
        doc.paragraphs[1].text = "Короткий заголовок"
        doc.paragraphs[7].text = "Аннотация. Слишком коротко."
        doc.paragraphs[8].text = "Ключевые слова: PLM; слишком длинная ключевая фраза из пяти слов"
        doc.paragraphs[15].text = "Abstract. Too short."
        doc.paragraphs[16].text = "Keywords: AI, one, two, three"

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("УДК должен быть указан в формате «УДК 520.607»", errors_ru)
        self.assertIn("Заголовок статьи должен содержать от 6 до 15 слов", errors_ru)
        self.assertIn("Аннотация должна содержать от 650 до 1000 знаков с пробелами", errors_ru)
        self.assertIn("Аннотация на английском языке должна содержать от 650 до 1000 знаков с пробелами", errors_ru)
        self.assertIn("Ключевые слова должны разделяться запятыми", errors_ru)
        self.assertIn("Ключевые слова не должны содержать аббревиатуры", errors_ru)
        self.assertIn("Ключевых слов на английском языке должно быть от 5 до 7", errors_ru)

    def test_keywords_must_not_end_with_period(self):
        doc = _valid_article_document()
        doc.paragraphs[8].runs[-1].text = (
            " энергообеспечение, загрязнение окружающей среды, нанотехнологии, "
            "экологическое равновесие, свойства наноматериалов."
        )
        doc.paragraphs[16].runs[-1].text = (
            " energy supply, environmental pollution, nanotechnology, ecological balance, material properties."
        )

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Ключевые слова не должны заканчиваться точкой", errors_ru)
        self.assertIn("Ключевые слова на английском языке не должны заканчиваться точкой", errors_ru)

    def test_structural_formatting_violations_are_reported(self):
        doc = _valid_article_document()
        doc.paragraphs[1].runs[0].bold = False
        doc.paragraphs[5].runs[0].italic = False
        doc.paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.LEFT
        _main_text_paragraph(doc).paragraph_format.first_line_indent = Cm(0.5)
        for paragraph in doc.paragraphs:
            if paragraph.text == "References":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                break

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn(
            "Заголовок статьи должен быть выделен полужирным и не должен быть набран курсивом",
            errors_ru,
        )
        self.assertIn("Аффилиации должны быть набраны курсивом без полужирного начертания", errors_ru)
        self.assertIn("Аннотация должна быть выровнена по ширине", errors_ru)
        self.assertIn("Абзацный отступ основного текста должен быть 0, 1 см или 1.25 см", errors_ru)
        self.assertIn("Заголовок References должен быть выровнен по левому краю или по ширине", errors_ru)

    def test_main_text_must_not_mix_one_and_one_twenty_five_cm_indents(self):
        doc = _valid_article_document()
        _main_text_paragraph(doc).paragraph_format.first_line_indent = Cm(1)
        second_main_paragraph = _insert_main_text_paragraph_before_references(
            doc,
            "A second main-text paragraph keeps the same article body before references.",
        )
        second_main_paragraph.paragraph_format.first_line_indent = Cm(1.25)

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn(
            "В документе нельзя одновременно использовать абзацные отступы 1 см и 1.25 см",
            errors_ru,
        )

    def test_tables_figures_and_formula_objects_pass(self):
        doc = _valid_article_document()
        _add_table_caption_and_title(doc, 1)
        _add_table(doc)
        _add_small_picture(doc)
        _add_article_paragraph(doc, "Рисунок 1 Example figure")

        paragraph = _add_article_paragraph(doc, "Formula: ")
        formula_run = paragraph.add_run("x")
        formula_run.font.name = "Cambria Math"
        formula_run.font.size = Pt(12)
        formula_run._element.append(
            parse_xml(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                "<m:r><m:t>x</m:t></m:r>"
                "</m:oMath>"
            )
        )

        errors_ru, errors_en = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertEqual([], errors_ru)
        self.assertEqual([], errors_en)

    def test_table_violations_are_reported(self):
        doc = _valid_article_document()
        for number in (1, 2, 3):
            _add_table_caption_and_title(doc, number)
            _add_table(doc)

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("В статье должно быть не более 2 таблиц", errors_ru)

    def test_table_formatting_violations_are_reported(self):
        doc = _valid_article_document()
        _add_article_paragraph(doc, "Таблица 1")
        _add_article_paragraph(doc, "Заголовок таблицы")
        table = doc.add_table(rows=1, cols=1)
        paragraph = table.rows[0].cells[0].paragraphs[0]
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run("1")
        run.font.name = "Arial"
        run.font.size = Pt(10)

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Номер таблицы должен быть выровнен по правому краю и набран курсивом", errors_ru)
        self.assertIn("Заголовок таблицы должен быть выровнен по центру и выделен полужирным", errors_ru)
        self.assertIn("В таблицах необходимо использовать шрифт Times New Roman", errors_ru)

    def test_figure_violations_are_reported(self):
        doc = _valid_article_document()
        _add_small_picture(doc)

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Все рисунки должны иметь подрисуночные подписи", errors_ru)

    def test_caption_sequence_and_plain_text_formula_violations_are_reported(self):
        doc = _valid_article_document()
        _add_table_caption_and_title(doc, 2)
        _add_table(doc)
        _add_article_paragraph(doc, "E = mc2")

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Таблицы должны нумероваться последовательно в порядке упоминания", errors_ru)
        self.assertIn("Формулы должны быть набраны в редакторе формул Word, Equation или MathType", errors_ru)

    def test_reference_format_order_and_missing_sources_are_reported(self):
        doc = _valid_article_document()
        _main_text_paragraph(doc).text = (
            "Introduction. First mention is out of order [2]. "
            "Then the first source appears [1, с. 12]. "
            "This cites a missing source [6, p. 7] and contains a bad marker [bad]. "
        )

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Неверный формат ссылки: [bad]", errors_ru)
        self.assertIn("В тексте есть ссылки на источники, отсутствующие в списке", errors_ru)
        self.assertIn("В тексте должны быть ссылки на все источники из списка", errors_ru)
        self.assertIn("Список источников должен формироваться в порядке первого упоминания в тексте", errors_ru)

    def test_reference_formatting_violations_are_reported(self):
        doc = _valid_article_document()
        first_reference = _reference_paragraphs(doc)[0]
        first_reference.text = "[1] Source title 1. Moscow, Publisher, 2024, 10 p."
        first_reference.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first_reference.paragraph_format.left_indent = Cm(0)
        first_reference.paragraph_format.first_line_indent = Cm(0)
        first_reference.runs[0].font.name = "Times New Roman"
        first_reference.runs[0].font.size = Pt(12)
        first_reference.runs[0].bold = True

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn(
            "Элементы списка источников должны начинаться с формата «[N] » и табуляции после пробела",
            errors_ru,
        )
        self.assertIn("Элементы списка источников должны быть выровнены по левому краю или по ширине", errors_ru)
        self.assertIn("Элементы списка источников должны иметь висячий отступ, равный левому отступу", errors_ru)
        self.assertIn("Элементы списка источников не должны содержать полужирное начертание", errors_ru)

    def test_reference_count_and_numbering_violations_are_reported(self):
        doc = _valid_article_document()
        reference_paragraphs = _reference_paragraphs(doc)
        for paragraph in reference_paragraphs[2:]:
            paragraph.clear()
        reference_paragraphs[1].text = "[4] Source title 4. Moscow, Publisher, 2024, 10 p."

        errors_ru, _ = Validator(_docx_bytes(doc, pages=4)).validate()

        self.assertIn("Список источников должен содержать не менее 5 источников", errors_ru)
        self.assertIn("Источники в списке должны быть пронумерованы последовательно, начиная с [1]", errors_ru)


if __name__ == "__main__":
    unittest.main()
