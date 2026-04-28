from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from i18n import (
    DEFAULT_LANGUAGE,
    field_label,
    participation_label,
    participation_status_label,
    review_status_label,
    section_label,
    text,
    validation_status_label,
    validation_summary_label,
)
from models import PARTICIPATION_STATUSES, REVIEW_STATUSES


MOSCOW_TZ = timezone(timedelta(hours=3), name="UTC+3")

APPLICATION_EXPORT_FIELDS: tuple[str, ...] = (
    "_id",
    "owner_email",
    "form_language",
    "last_name",
    "first_name",
    "middle_name",
    "place_of_study",
    "department",
    "place_of_work",
    "job_title",
    "phone",
    "email",
    "participation",
    "section",
    "publication_title",
    "foreign_language_consultant",
    "participation_status",
    "review_status",
    "publication_file.filename",
    "expert_opinion_file.filename",
    "review_file.filename",
    "publication_validation.status",
    "publication_validation.summary",
    "publication_validation.errors",
    "publication_validation.last_error",
    "publication_validation.updated_at",
    "comments",
    "created_at",
    "updated_at",
)


def build_applications_xlsx(records: list[dict[str, Any]], *, lang: str = DEFAULT_LANGUAGE) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Applications"

    worksheet.append([field_label(lang, field_path) for field_path in APPLICATION_EXPORT_FIELDS])
    for record in records:
        worksheet.append(
            [_export_value(record, field_path, lang=lang) for field_path in APPLICATION_EXPORT_FIELDS]
        )

    _style_worksheet(worksheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_worksheet(worksheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="0F5959")
    header_font = Font(bold=True, color="FFFFFF")
    thin_line = Side(style="thin", color="D7D2C8")
    cell_border = Border(left=thin_line, right=thin_line, top=thin_line, bottom=thin_line)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = cell_border

    for row in worksheet.iter_rows(min_row=2):
        worksheet.row_dimensions[row[0].row].height = 18
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = cell_border

    worksheet.row_dimensions[1].height = 24
    for column_index in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = _column_width(column_index)


def _column_width(index: int) -> int:
    if index in {1, 2, 19, 20, 21}:
        return 24
    if index in {13, 14, 15, 16, 22, 23, 24, 27}:
        return 34
    return 20


def _export_value(record: dict[str, Any], field_path: str, *, lang: str) -> str:
    if field_path == "_id":
        return str(record.get("_id") or "")
    if field_path == "owner_user_id":
        return str(record.get("owner_user_id") or "")
    if field_path == "form_language":
        return _form_language_label(record.get("form_language"), lang=lang)
    if field_path == "participation":
        return participation_label(lang, str(record.get("participation") or ""))
    if field_path == "section":
        return section_label(lang, str(record.get("section") or ""))
    if field_path == "participation_status":
        return participation_status_label(
            lang,
            str(record.get("participation_status") or PARTICIPATION_STATUSES[0]),
        )
    if field_path == "review_status":
        return review_status_label(lang, str(record.get("review_status") or REVIEW_STATUSES[0]))
    if field_path == "comments":
        return _comments_text(record.get("comments"), lang=lang)
    if field_path.startswith("publication_validation."):
        return _publication_validation_value(record, field_path, lang=lang)
    if "." in field_path:
        root_key, nested_key = field_path.split(".", 1)
        nested = record.get(root_key)
        if not isinstance(nested, dict):
            return ""
        return _string_value(nested.get(nested_key), lang=lang)
    return _string_value(record.get(field_path), lang=lang)


def _publication_validation_value(record: dict[str, Any], field_path: str, *, lang: str) -> str:
    validation_info = record.get("publication_validation")
    nested_key = field_path.split(".", 1)[1]
    if not isinstance(validation_info, dict):
        return ""

    nested_value = validation_info.get(nested_key)
    if nested_key == "status":
        return validation_status_label(lang, str(nested_value or ""))
    if nested_key == "summary":
        return validation_summary_label(lang, str(nested_value or ""))
    if nested_key == "errors":
        if not isinstance(nested_value, list):
            return ""
        return "\n".join(str(item).strip() for item in nested_value if str(item).strip())
    if nested_key == "updated_at":
        return _format_dt(nested_value, lang=lang)
    return str(nested_value or "")


def _string_value(value: Any, *, lang: str) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return _format_dt(value, lang=lang)
    if isinstance(value, (list, tuple)):
        return _clean_cell_text("\n".join(str(item).strip() for item in value if str(item).strip()))
    return _clean_cell_text(str(value))


def _format_dt(value: Any, *, lang: str) -> str:
    if not isinstance(value, datetime):
        return ""
    localized = value.astimezone(MOSCOW_TZ)
    suffix = text(lang, "timezone_suffix")
    return f"{localized:%d.%m.%Y %H:%M} {suffix}"


def _form_language_label(value: Any, *, lang: str) -> str:
    value_text = str(value or "").strip().lower()
    if value_text == "ru":
        return text(lang, "language_ru")
    if value_text == "en":
        return text(lang, "language_en")
    return _clean_cell_text(value_text)


def _comments_text(value: Any, *, lang: str) -> str:
    if not isinstance(value, list):
        return ""

    parts: list[str] = []
    for item in value:
        if not isinstance(item, dict):
            continue
        author_role = str(item.get("author_role") or "").strip().lower()
        if author_role == "admin":
            author = text(lang, "comment_author_admin")
        elif author_role == "author":
            author = text(lang, "comment_author_author")
        else:
            author = text(lang, "comment_author_unknown")
        author_email = str(item.get("author_email") or "").strip()
        created_at = _format_dt(item.get("created_at"), lang=lang)
        comment = str(item.get("text") or "").strip()
        meta = " | ".join(part for part in (author, author_email, created_at) if part)
        parts.append(f"{meta}: {comment}" if comment else meta)
    return _clean_cell_text("\n".join(parts))


def _clean_cell_text(value: str) -> str:
    cleaned = ILLEGAL_CHARACTERS_RE.sub(" ", value)
    if cleaned.startswith(("=", "+", "-", "@")):
        return f"'{cleaned}"
    return cleaned
